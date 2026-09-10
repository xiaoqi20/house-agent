"""解析层测试：只用 tests/samples/ 里的固定样本，离线且确定性（不碰网络/LLM）。"""

import io
import time
from pathlib import Path

import pytest

from rentgraph.services import parsing as P

SAMPLES = Path(__file__).parent / "samples"


def _read(name: str) -> bytes:
    return (SAMPLES / name).read_bytes()


def test_supported_extensions_frozen() -> None:
    assert P.SUPPORTED_EXTENSIONS == frozenset({".txt", ".csv", ".md", ".xlsx", ".xls", ".docx", ".pdf"})


def test_parse_txt_keeps_listings_and_newlines() -> None:
    doc = P.parse_document("houses.txt", _read("houses.txt"))
    assert doc.kind == "txt"
    assert doc.rows is None
    assert doc.page_count == 1
    assert doc.pages == []
    for name in ("望京花园", "望京西园三区", "中关村南大街 40 号院", "天通苑北一区"):
        assert name in doc.text
    assert "5800" in doc.text
    # 单个换行必须保留（合同条款切分依赖原文换行），只折叠 3+ 连续空行
    assert "\n" in doc.text


def test_parse_csv_utf8() -> None:
    doc = P.parse_document("houses.csv", _read("houses.csv"))
    assert doc.kind == "csv"
    # CSV 与 Excel 走同一套表格映射（W5 修复：表头 + 数据行 → 「表头: 值」），不再当纯文本直读
    assert doc.rows is not None and len(doc.rows) == 5  # 1 表头 + 4 房源
    assert doc.rows[1][0] == "望京花园"
    assert doc.rows[1][4] == "5800"
    assert doc.text == P.rows_to_text(doc.rows)
    assert "小区: 望京花园" in doc.text and "租金: 5800" in doc.text


def test_parse_csv_gbk_fallback() -> None:
    raw = "小区,租金\n望京花园,5800\n".encode("gbk")
    with pytest.raises(UnicodeDecodeError):
        raw.decode("utf-8")  # 前提校验：确实不是 utf-8，才走 gb18030 兜底
    doc = P.parse_document("houses.csv", raw)
    assert "望京花园" in doc.text
    assert "\ufffd" not in doc.text


def test_parse_md() -> None:
    doc = P.parse_document("houses.md", _read("houses.md"))
    assert doc.kind == "md"
    assert "## 1. 望京花园" in doc.text
    assert "3200 元/月" in doc.text


def test_parse_xlsx_preserves_rows_as_strings() -> None:
    doc = P.parse_document("houses.xlsx", _read("houses.xlsx"))
    assert doc.kind == "xlsx"
    assert doc.rows is not None and len(doc.rows) == 5  # 1 表头 + 4 房源
    assert all(isinstance(c, str) for row in doc.rows for c in row)
    assert doc.rows[0][0] == "小区"
    assert doc.rows[1][4] == "5800"  # 数字单元格转字符串，不带 .0
    assert doc.text == P.rows_to_text(doc.rows)
    assert "小区: 望京花园" in doc.text and "租金: 5800" in doc.text


def test_parse_xls_preserves_rows_as_strings() -> None:
    doc = P.parse_document("houses.xls", _read("houses.xls"))
    assert doc.kind == "xls"
    assert doc.rows is not None and len(doc.rows) == 5
    assert doc.rows[1][0] == "望京花园"
    assert doc.rows[1][4] == "5800"
    assert doc.text == P.rows_to_text(doc.rows)


def test_parse_xlsx_generated_on_the_fly(tmp_path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["小区", "租金"])
    ws.append(["望京花园", 5800])
    p = tmp_path / "gen.xlsx"
    wb.save(p)
    doc = P.parse_document("gen.xlsx", p.read_bytes())
    assert doc.rows == [["小区", "租金"], ["望京花园", "5800"]]
    assert doc.text == "小区: 望京花园 | 租金: 5800"


def test_parse_docx_paragraphs_and_tables() -> None:
    import docx

    d = docx.Document()
    d.add_paragraph("望京花园 2室1厅 60平")
    d.add_paragraph("")  # 空段落应被丢弃
    d.add_paragraph("月租 5800 元，押一付三")
    table = d.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "费用项目"
    table.rows[0].cells[1].text = "金额"
    table.rows[1].cells[0].text = "物业费"
    table.rows[1].cells[1].text = "300"
    buf = io.BytesIO()
    d.save(buf)

    doc = P.parse_document("houses.docx", buf.getvalue())
    assert doc.kind == "docx"
    assert doc.text.startswith("望京花园 2室1厅 60平\n月租 5800 元，押一付三")
    assert "费用项目: 物业费 | 金额: 300" in doc.text  # 表格按「表头: 值」并入正文


def test_parse_pdf_text_layer_pages_and_clauses() -> None:
    import re

    doc = P.parse_document("contract.pdf", _read("contract.pdf"))
    assert doc.kind == "pdf"
    assert doc.page_count == 2  # 真实页数（含分页符的固定样本）
    assert [p.number for p in doc.pages] == [1, 2]
    assert doc.pages[0].text and doc.pages[1].text
    clauses = re.findall(r"第[一二三四五六七八九十]+条", doc.text)
    assert "第一条" in clauses and "第十二条" in clauses
    assert "押金" in doc.text


def test_scanned_pdf_raises_no_text_layer() -> None:
    with pytest.raises(P.NoTextLayer) as exc:
        P.parse_document("scanned.pdf", _read("scanned.pdf"))
    assert exc.value.code == "NO_TEXT_LAYER"
    assert "扫描件" in exc.value.message
    assert "0 个字符" in exc.value.hint  # 字符数写进 hint，供前端提示与 API 阈值判断


def test_empty_file_raises_empty_document() -> None:
    with pytest.raises(P.EmptyDocument) as exc:
        P.parse_document("empty.txt", b"")
    assert exc.value.code == "EMPTY_DOCUMENT"
    with pytest.raises(P.EmptyDocument):
        P.parse_document("blank.txt", b"   \n\t\n  ")


def test_doc_returns_doc_unsupported() -> None:
    with pytest.raises(P.UnsupportedFormat) as exc:
        P.parse_document("合同.doc", b"\xd0\xcf\x11\xe0binary")
    assert exc.value.code == "DOC_UNSUPPORTED"
    assert "docx" in exc.value.hint


def test_unsupported_extension_lists_supported_formats() -> None:
    with pytest.raises(P.UnsupportedFormat) as exc:
        P.parse_document("deck.pptx", b"x")
    assert exc.value.code == "UNSUPPORTED_FORMAT"
    for ext in P.SUPPORTED_EXTENSIONS:
        assert ext in exc.value.hint


def test_corrupt_xlsx_is_typed_error_not_silent() -> None:
    with pytest.raises(P.ParseError) as exc:
        P.parse_document("broken.xlsx", b"definitely not a workbook")
    assert exc.value.code == "PARSE_FAILED"


def test_parse_text_paste_path() -> None:
    doc = P.parse_text("房源 A\u00a0月租 5800 元\n\n\n\n房源 B 月租 5200 元")
    assert doc.kind == "text"
    assert doc.pages == []
    assert doc.rows is None
    assert doc.page_count == 1
    assert "\u00a0" not in doc.text  # NBSP 归一
    assert "\n\n\n" not in doc.text  # 空行折叠
    assert "房源 A 月租 5800 元" in doc.text


def test_parse_text_rejects_blank_paste() -> None:
    with pytest.raises(P.EmptyDocument):
        P.parse_text("   \n  ")


def test_rows_to_text_uses_header_or_generates_columns() -> None:
    with_header = [["小区", "租金", "押金"], ["望京花园", "5800", "押一付三"]]
    assert P.rows_to_text(with_header) == "小区: 望京花园 | 租金: 5800 | 押金: 押一付三"
    # 首行像数据（含金额）时不当表头，退回 列1..列N，避免吃掉第一套房源
    no_header = [["望京花园", "5800"], ["望京西园三区", "5200"]]
    assert P.rows_to_text(no_header) == "列1: 望京花园 | 列2: 5800\n列1: 望京西园三区 | 列2: 5200"
    assert P.rows_to_text([]) == ""
    assert P.rows_to_text([["", " "]]) == ""


def test_large_text_file_parses_fast() -> None:
    chunk = "望京花园 2室1厅 60平 月租 5800 元 押一付三 距地铁 600 米\n"
    unit = len(chunk.encode())
    data = (chunk * (-(-4 * 1024 * 1024 // unit))).encode()
    assert len(data) >= 4 * 1024 * 1024
    start = time.perf_counter()
    doc = P.parse_document("big.txt", data)
    elapsed = time.perf_counter() - start
    assert doc.kind == "txt" and doc.text
    assert elapsed < 1.5, f"4MB 文本解析耗时 {elapsed:.3f}s，超出预算"
